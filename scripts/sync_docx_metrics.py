import os
import sys

# Configure stdout to handle UTF-8 printing on Windows
if sys.stdout.encoding.lower() != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')

try:
    from docx import Document
except ImportError:
    print("Thư viện 'python-docx' chưa được cài đặt. Hãy chạy lệnh: pip install python-docx")
    sys.exit(1)

# Từ điển ánh xạ các giá trị cũ sang giá trị mới (Nguồn chân lý duy nhất)
REPLACEMENTS = {
    # Dataset counts
    "34.054": "34.168",
    "27.243": "27.334",
    "6.811": "6.834",
    "27.219": "27.310",
    
    # Anomaly injection counts
    "204 mẫu": "205 mẫu",
    "545": "546",
    
    # Hyper-parameters
    "n_estimators=100": "n_estimators=200",
    "n_estimators = 100": "n_estimators = 200",
    "max_samples=256": "max_samples=512",
    "max_samples = 256": "max_samples = 512",
    
    # Storage
    "4,3 MB": "4,36 MB",
    
    # Performance Metrics
    "94,1%": "61,95%",
    "92,7%": "94,15%",
    "83,6%": "57,04%",
    
    # Inference Latency
    "≈ 1,2 ms": "≈ 18-26 ms",
    "1,2 ms": "18-26 ms"
}

def replace_in_runs(paragraph):
    """
    Duyệt qua các 'run' trong một đoạn văn (paragraph) và thay thế văn bản.
    Việc thay thế trên 'run' giúp giữ nguyên định dạng in đậm, in nghiêng...
    """
    for key, val in REPLACEMENTS.items():
        if key in paragraph.text:
            # Thay thế trên từng run
            for run in paragraph.runs:
                if key in run.text:
                    run.text = run.text.replace(key, val)
            
            # Nếu giá trị vẫn còn trong paragraph.text, có thể chuỗi bị chia cắt giữa nhiều run
            # Xử lý bằng cách ghi đè lại toàn bộ đoạn văn (sẽ mất định dạng nội bộ của đoạn đó)
            if key in paragraph.text:
                paragraph.text = paragraph.text.replace(key, val)

def process_document(input_path, output_path):
    print(f"Đang mở tệp: {input_path}")
    doc = Document(input_path)
    
    print("Đang xử lý các đoạn văn bản chính...")
    for para in doc.paragraphs:
        replace_in_runs(para)
        
    print("Đang xử lý các bảng biểu...")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    replace_in_runs(para)
                    
    print(f"Đang lưu tệp đã cập nhật tại: {output_path}")
    doc.save(output_path)
    print("Thành công!")

if __name__ == "__main__":
    # Xác định đường dẫn gốc của project (thư mục chứa thư mục scripts và docs)
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(script_dir)
    
    docs_dir = os.path.join(project_root, "docs")
    input_docx = os.path.join(docs_dir, "Bao_Cao_Do_An_Nganh_Hoan_Chinh.docx")
    output_docx = os.path.join(docs_dir, "Bao_Cao_Do_An_Nganh_Hoan_Chinh_Synced.docx")
    
    if not os.path.exists(input_docx):
        print(f"Lỗi: Không tìm thấy tệp {input_docx}")
        sys.exit(1)
        
    process_document(input_docx, output_docx)
